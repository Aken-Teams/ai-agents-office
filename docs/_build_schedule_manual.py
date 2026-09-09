# -*- coding: utf-8 -*-
"""Build 排程功能操作手冊 (.docx).

Layout follows the project-management-system 《系統操作手冊 - v01》 design:
navy headings, "● 步驟 N" run-ins, framed screenshot placeholders, and a
section rule between chapters. Scope is the team-schedule feature only, so the
manual stays short.

    python docs/_build_schedule_manual.py
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CJK   = "Microsoft JhengHei"
NAVY  = RGBColor(0x1E, 0x3A, 0x8A)
BLUE  = RGBColor(0x25, 0x63, 0xEB)
INK   = RGBColor(0x0F, 0x17, 0x2A)
SUB   = RGBColor(0x33, 0x41, 0x55)
MUTE  = RGBColor(0x64, 0x74, 0x8B)
RED   = RGBColor(0xB9, 0x1C, 0x1C)

CONTENT_W = Cm(16.0)
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO = os.path.join(ROOT, 'client', 'public', 'logo.png')
OUT  = os.path.join(ROOT, 'docs', '排程功能操作手冊 - v01.docx')

doc = Document()
normal = doc.styles['Normal']
normal.font.name = CJK
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), CJK)
normal.paragraph_format.line_spacing = 1.35
for s in doc.sections:
    s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.5); s.right_margin = Cm(2.5)


# -- low-level helpers -------------------------------------------------------
def _cjk(run):
    run.font.name = CJK
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), CJK)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def borders(cell, color="D8E0EA", sz="6"):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), sz)
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        b.append(e)
    tcPr.append(b)


def set_w(cell, w):
    cell.width = w
    tcPr = cell._tc.get_or_add_tcPr()
    tw = OxmlElement('w:tcW')
    tw.set(qn('w:w'), str(int(w.twips))); tw.set(qn('w:type'), 'dxa')
    tcPr.append(tw)


def para(text="", size=11, color=SUB, bold=False, align=None,
         before=2, after=5, indent=None, lh=1.35):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = lh
    if align:
        p.alignment = align
    if indent is not None:
        pf.left_indent = indent
    if text:
        r = p.add_run(text); r.bold = bold
        r.font.size = Pt(size); r.font.color.rgb = color; _cjk(r)
    return p


def h1(text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(20); pf.space_after = Pt(6)
    pf.keep_with_next = True
    r = p.add_run(text); r.bold = True
    r.font.size = Pt(16); r.font.color.rgb = NAVY; _cjk(r)
    return p


def step(label, title):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(12); pf.space_after = Pt(2)
    pf.keep_with_next = True
    r = p.add_run("● "); r.bold = True
    r.font.size = Pt(12); r.font.color.rgb = BLUE; _cjk(r)
    r2 = p.add_run(label + "　" + title); r2.bold = True
    r2.font.size = Pt(12); r2.font.color.rgb = NAVY; _cjk(r2)
    return p


def bullet(text, size=10.5):
    p = doc.add_paragraph(style='List Bullet')
    pf = p.paragraph_format
    pf.space_before = Pt(1); pf.space_after = Pt(1); pf.line_spacing = 1.3
    r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = SUB; _cjk(r)
    return p


def shot(caption, height=Cm(4.6)):
    """Empty framed box for the reader to paste the real screenshot into.

    The height comes from the row rule (an empty paragraph has none), so the
    frame stands open with nothing to delete before pasting a screenshot.
    """
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    row = t.rows[0]
    row.height = height
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell = t.cell(0, 0)
    borders(cell, "BFDBFE"); set_w(cell, CONTENT_W)
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(3); cap.paragraph_format.space_after = Pt(10)
    rc = cap.add_run(caption)
    rc.font.size = Pt(9.5); rc.italic = True; rc.font.color.rgb = MUTE; _cjk(rc)


def note(label, text, fill="FFF7ED", edge="FDBA74", labelcolor=RGBColor(0xB4, 0x53, 0x09)):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    shade(cell, fill); borders(cell, edge); set_w(cell, CONTENT_W)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    r1 = p.add_run("  " + label + "　"); r1.bold = True
    r1.font.size = Pt(10.5); r1.font.color.rgb = labelcolor; _cjk(r1)
    r2 = p.add_run(text); r2.font.size = Pt(10.5); r2.font.color.rgb = INK; _cjk(r2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        cell = t.rows[0].cells[i]
        shade(cell, "E8EEF7"); borders(cell); set_w(cell, widths[i])
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
        r = p.add_run(htxt); r.bold = True
        r.font.size = Pt(10.5); r.font.color.rgb = NAVY; _cjk(r)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cell = cells[i]
            borders(cell); set_w(cell, widths[i])
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.25
            r = p.add_run(val); r.font.size = Pt(10.5)
            r.font.color.rgb = INK if i == 0 else SUB
            r.bold = (i == 0)
            _cjk(r)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def rule():
    """Thin separator closing a chapter."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(0)
    pr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '4'); bottom.set(qn('w:color'), 'D8E0EA')
    pbdr.append(bottom); pr.append(pbdr)
    r = p.add_run(""); r.font.size = Pt(2); _cjk(r)


def page_break():
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


# -- Cover -------------------------------------------------------------------
for _ in range(3):
    doc.add_paragraph()
if os.path.exists(LOGO):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(LOGO, width=Cm(3.6))
para("AI Agents Office", size=32, color=INK, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=2)
para("排程功能操作手冊", size=19, color=NAVY, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
_rule = doc.add_paragraph(); _rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
_pr = _rule._p.get_or_add_pPr(); _pbdr = OxmlElement('w:pBdr')
_b = OxmlElement('w:bottom')
_b.set(qn('w:val'), 'single'); _b.set(qn('w:sz'), '12')
_b.set(qn('w:space'), '2'); _b.set(qn('w:color'), '1E3A8A')
_pbdr.append(_b); _pr.append(_pbdr)
_r = _rule.add_run("　" * 24); _r.font.size = Pt(2); _cjk(_r)
para("版本 v1.0　|　日期 2026-08-30　|　適用對象：使用 AI 團隊定期報告的成員",
     size=10.5, color=MUTE, align=WD_ALIGN_PARAGRAPH.CENTER, before=14, after=4)
para("本手冊僅涵蓋「排程」功能：如何建立定期執行的 AI 團隊分析、"
     "設定收件者與產出文件，以及如何確認每次是否如期執行。",
     size=10.5, color=SUB, align=WD_ALIGN_PARAGRAPH.CENTER, after=6, lh=1.45)

page_break()

# -- 目錄 --------------------------------------------------------------------
para("目　錄", size=16, color=NAVY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
for line in [
    "如何使用本手冊",
    "一、排程能做什麼",
    "二、進入排程管理頁",
    "三、新增排程（四個步驟）",
    "四、管理既有排程",
    "五、確認是否如期執行",
    "六、收到報告之後",
    "七、常見問題",
]:
    para(line, size=11, color=INK, indent=Cm(1.2), before=1, after=4)

# -- 如何使用本手冊 -----------------------------------------------------------
h1("如何使用本手冊")
para("本手冊以操作任務分章，逐步說明每個步驟與預期畫面；每個步驟下方留有空白框，"
     "請將實際操作畫面截圖貼入框內後再發送給使用者。排程只是「定時自動執行 AI 團隊分析」，"
     "團隊本身怎麼建立、怎麼即時協作，請參閱系統內的「使用指南」。")
note("先看這裡",
     "排程一律建立在某一個「AI 團隊」底下：時間到了，系統會用你設定的議題文字自動重跑一次該團隊的分析，"
     "把結果寄到指定信箱（可另外附上 Word／簡報／PDF），並把每一次執行記錄在排程管理頁。",
     fill="EFF6FF", edge="BFDBFE", labelcolor=NAVY)
rule()

# -- 一、排程能做什麼 ---------------------------------------------------------
h1("一、排程能做什麼")
para("建立排程前，先確認排程的能力範圍，避免設定後拿不到預期的結果。")
table(
    ["項目", "說明"],
    [
        ["執行內容", "以你填寫的「分析議題」文字，重新執行該 AI 團隊的協作分析。"],
        ["執行頻率", "每天，或每週指定星期幾；可指定到分鐘（例如 08:30）。"],
        ["產出", "分析報告（信件內文）；可另外選擇產生簡報、Word、PDF 或網頁簡報，以下載連結附在信中。"],
        ["送達方式", "寄送 E-mail，可設定多位收件者；未來將整合強茂內部 IM 系統推播。"],
        ["執行位置", "伺服器背景執行，不需要開著瀏覽器，也看不到逐字的即時過程。"],
        ["不支援", "排程無法帶入你上傳的檔案，每次都是重新分析議題文字。"],
    ],
    [Cm(3.4), Cm(12.6)],
)
note("排程不會帶檔案",
     "排程只保存「議題文字」，不會保存你上傳的檔案，因此需要依據上傳資料分析的題目請改用即時分析"
     "（回到團隊頁上傳檔案後執行）。從有上傳檔案的分析按「排程」時，系統會擋下並提示。"
     "排程適合盤勢、新聞、公開資訊等會持續更新的議題。")
rule()

# -- 二、進入排程管理頁 -------------------------------------------------------
h1("二、進入排程管理頁")
step("步驟 1", "從左側「AI 助手」開啟團隊")
para("登入後點左側選單「AI 助手」，在團隊區塊點要排程的團隊，進入該團隊的協作頁。")
shot("圖 2-1：AI 助手 — 團隊清單")

step("步驟 2", "點右上角「排程」")
para("在團隊頁右上角工具列點「排程」，進入該團隊的排程管理頁。"
     "若剛在輸入框打好議題才點進來，系統會自動帶入該議題並開啟新增視窗，不必重打。")
shot("圖 2-2：團隊頁 — 右上角「排程」按鈕")

step("步驟 3", "認識排程管理頁")
para("排程管理頁分為三個區塊，手機版以上方分頁切換：")
bullet("排程清單：這個團隊的所有排程，以及各自的下次執行時間。")
bullet("執行日曆：以月曆呈現哪幾天實際跑過、哪一天有預定要跑。")
bullet("執行紀錄：每一次執行的時間、結果與是否寄送成功。")
shot("圖 2-3：排程管理頁 — 清單／日曆／執行紀錄")
note("權限",
     "排程功能目前僅開放「管理者」與「檢閱者」使用；一般使用者不會看到排程按鈕，"
     "直接開啟排程網址也會被導回團隊頁。")
rule()


# -- 三、新增排程 -------------------------------------------------------------
h1("三、新增排程（四個步驟）")
para("按右上角「新增排程」開啟設定視窗，依序完成四個步驟：基本 → 週期與文件 → 收件者 → 確認。"
     "未填必填欄位時，「下一步」會維持不可按的狀態。")

step("步驟 1", "填寫排程名稱與分析議題")
para("「排程名稱」是給人看的識別名稱（例如：每日台股盤勢），會出現在清單、執行紀錄與信件標題；"
     "「分析議題」才是團隊每次真正要分析的內容，寫得越具體，產出的報告越聚焦。兩者皆為必填。")
shot("圖 3-1：步驟 1 — 排程名稱與分析議題")

step("步驟 2", "設定執行週期與產出文件")
para("先選「每天」或「每週」；選每週時再指定星期幾，並設定執行時間（24 小時制）。")
para("下方可選擇是否「產出報告文件」：不產生（預設，只寄文字報告）、簡報、Word、PDF 或網頁簡報。"
     "選定格式後，再挑一個版面風格。")
shot("圖 3-2：步驟 2 — 週期、時間與文件格式")
note("產檔會比較久",
     "產生文件在背景進行（簡報約 3–8 分鐘），不會影響報告信的寄送；完成後下載連結會附在信件中。",
     fill="EFF6FF", edge="BFDBFE", labelcolor=NAVY)

step("步驟 3", "加入收件者")
para("先選公司／網域（台灣 PANJIT、璟茂、無錫強茂、強茂深圳、蘇州群鑫、強茂徐州、山東強茂），"
     "再用搜尋框以姓名或帳號找人，點一下即加入收件清單，系統會自動帶出該員的公司信箱。"
     "可加入多位，報告會同時寄給所有人；點名條右側的「×」可移除。")
shot("圖 3-3：步驟 3 — 從 AD 目錄挑選收件者")

step("步驟 4", "確認並建立")
para("最後一頁列出排程名稱、分析議題、執行週期、產出文件與收件者名單；"
     "任何一列都可以點「編輯」跳回該步驟修改。確認無誤後按「建立排程」。"
     "建立完成後回到清單，該筆排程即顯示下次執行時間。")
shot("圖 3-4：步驟 4 — 設定確認頁")
rule()


# -- 四、管理既有排程 ---------------------------------------------------------
h1("四、管理既有排程")
para("排程清單每一列右側都有操作按鈕，由左至右依序為：")
table(
    ["按鈕", "作用"],
    [
        ["立即測試執行", "馬上跑一次並寄出，用來驗證議題與收件設定；約 1 分鐘完成，不會更動原本的下次執行時間。"],
        ["編輯", "修改名稱、議題、週期、文件格式或收件者；儲存後會依新的時間重新計算下次執行。"],
        ["啟用／停用", "切換開關即可暫停排程；停用期間不會執行，重新啟用時會重新計算下次執行時間。"],
        ["刪除", "移除排程，需再次確認。刪除後不可復原，但已產生的執行紀錄與報告會保留。"],
    ],
    [Cm(3.4), Cm(12.6)],
)
shot("圖 4-1：排程清單 — 測試執行／編輯／啟用停用／刪除")
para("排程超過三筆時，清單上方會出現搜尋框，可用排程名稱、議題或收件者關鍵字快速篩選。")
rule()

# -- 五、確認是否如期執行 -----------------------------------------------------
h1("五、確認是否如期執行")
step("步驟 1", "看執行日曆")
para("日曆會標示每一天的狀況，顏色意義如下：")
table(
    ["標示", "意義"],
    [
        ["綠點＋數字", "當天實際執行成功的次數。"],
        ["紅點＋數字", "當天執行失敗或寄信失敗的次數。"],
        ["外框（預定 N）", "該日有 N 筆排程預定執行。"],
        ["整格底色", "今天。"],
    ],
    [Cm(3.4), Cm(12.6)],
)
para("點任一有標示的日期，可展開當天的「預定執行」與「執行紀錄」明細。")
shot("圖 5-1：執行日曆與當日明細")

step("步驟 2", "看執行紀錄")
para("頁面下方的執行紀錄依時間排列，右側標籤顯示每次執行的結果：")
bullet("背景執行中：正在跑，畫面會自動更新，不需手動重新整理。")
bullet("如期執行・已寄送：分析完成且信件寄出成功。")
bullet("如期執行・寄信失敗：分析成功但信沒寄出，通常是內部郵件閘道權限問題，請洽 IT。")
bullet("執行失敗：該次分析未完成，可等下次排程時間自動再跑，或按「立即測試執行」重跑。")
para("紀錄可用右上角下拉選單依排程篩選，並以分頁瀏覽；點右側的開新視窗圖示即可查看該次的完整報告。",
     before=4)
shot("圖 5-2：執行紀錄與狀態標籤")
rule()


# -- 六、收到報告之後 ---------------------------------------------------------
h1("六、收到報告之後")
para("排程完成後，報告會以下列方式送達：")
table(
    ["管道", "內容"],
    [
        ["E-mail", "主旨為【團隊協作報告】加上排程名稱；信中含分析議題、報告全文，以及「在網站上查看完整報告」按鈕；有選產出文件時另有下載按鈕。"],
        ["系統內", "團隊頁的「排程寄信」分頁，以及排程管理頁的執行紀錄，都可回頭查閱歷次報告。"],
        ["內部 IM（規劃中）", "未來將整合強茂內部 IM 系統，把每次報告的重點直接推播給你；目前以 E-mail 為送達管道。"],
    ],
    [Cm(3.6), Cm(12.4)],
)
shot("圖 6-1：報告信件與下載按鈕")
note("報告連結可分享",
     "信中的報告連結為公開連結，任何取得連結的人都能開啟，請斟酌轉寄對象。",
     fill="FEF2F2", edge="FECACA", labelcolor=RED)
rule()

# -- 七、常見問題 -------------------------------------------------------------
h1("七、常見問題")
table(
    ["問題", "說明"],
    [
        ["按了排程卻跳出「暫不支援此排程」", "該次分析有上傳檔案。排程無法帶入檔案，請改用即時分析，或把需要的內容寫進議題文字中。"],
        ["看不到「排程」按鈕", "排程目前僅開放管理者與檢閱者使用，請洽系統管理員調整權限。"],
        ["時間到了卻沒有執行", "系統每分鐘檢查一次，尖峰時可能略有延遲；伺服器若曾重啟，逾期的排程會在啟動後補跑。若當月用量已達上限，該次排程會略過不執行。"],
        ["顯示「寄信失敗」", "分析已完成（報告仍可在系統內查看），但郵件閘道未開放寄送權限，請洽 IT 協助。"],
        ["想暫停一陣子", "用清單上的開關停用即可，不需刪除；重新啟用後會從下一個週期繼續。"],
        ["測試執行會影響排程時間嗎", "不會。「立即測試執行」只是額外跑一次，原本的下次執行時間不變。"],
    ],
    [Cm(4.6), Cm(11.4)],
)
para("如仍有問題，請聯繫系統管理員，並提供排程名稱與該次執行時間，以利查詢紀錄。",
     size=10.5, color=MUTE, before=6)

doc.save(OUT)
print("saved:", OUT)
